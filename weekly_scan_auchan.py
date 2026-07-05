"""
Varredura semanal Auchan — detecta produtos novos e removidos nas secções conhecidas.

SETUP:
    pip install playwright pandas openpyxl
    python3 -m playwright install chromium

EXECUÇÃO:
    PYTHONUNBUFFERED=1 python3 weekly_scan_auchan.py
    PYTHONUNBUFFERED=1 python3 weekly_scan_auchan.py --test
    PYTHONUNBUFFERED=1 python3 weekly_scan_auchan.py --no-git

NOTAS:
    - Copia Scraping_Auchan.xlsx para Scraping_Auchan_YYYY-MM-DD.xlsx no início.
    - Varre as secções originais (scraping_auchan.py) + complemento (scraping_auchan_complemento.py).
    - Produtos já existentes: ignorados (sem visita à página de detalhe).
    - Produtos novos: extraídos e adicionados com letra azul.
    - Produtos removidos/indisponíveis: letra vermelha (linha mantida).
    - Checkpoint Excel + commit/push Git a cada 100 produtos novos processados.
    - Progresso: weekly_scan_auchan_progress.json
"""

from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import time
from datetime import date
from pathlib import Path

import pandas as pd
import scraping_auchan as auchan
from openpyxl import load_workbook
from openpyxl.styles import Font
from playwright.sync_api import Page, sync_playwright
from scraping_auchan_complemento import build_complement_targets, category_key

SOURCE_FILE = "Scraping_Auchan.xlsx"
PROGRESS_FILE = "weekly_scan_auchan_progress.json"
REPO_ROOT = Path(__file__).resolve().parent
GIT_PUSH_BACKOFF_SECONDS = (4, 8, 16, 32)
CHECKPOINT_EVERY = auchan.CHECKPOINT_EVERY

BLUE_FONT = Font(color="0000FF")
RED_FONT = Font(color="FF0000")

_git_checkpoint_enabled = True


def log(message: str) -> None:
    auchan.log(message)


def normalize_url(url: str) -> str:
    return auchan.normalize_product_url(str(url).strip())


def output_filename_for_date(scan_date: str) -> str:
    return f"Scraping_Auchan_{scan_date}.xlsx"


def build_weekly_scan_targets() -> tuple[auchan.CategoryTarget, ...]:
    """Secções originais + complemento, sem duplicar URLs de listagem."""
    targets: list[auchan.CategoryTarget] = list(auchan.CATEGORY_TARGETS)
    seen = {category_key(target.listing_url) for target in targets}
    for target in build_complement_targets():
        key = category_key(target.listing_url)
        if key not in seen:
            targets.append(target)
            seen.add(key)
    return tuple(targets)


def empty_progress(source_file: str, output_file: str, scan_date: str) -> dict:
    return {
        "scan_date": scan_date,
        "source_file": source_file,
        "output_file": output_file,
        "phase": "scanning",
        "completed_categories": [],
        "site_urls": [],
        "site_products": {},
        "baseline_urls": [],
        "new_urls_pending": [],
        "new_urls_processed": [],
        "removed_urls": [],
        "removed_urls_marked": [],
        "records": [],
        "new_products_added": 0,
        "in_progress": None,
    }


def load_progress() -> dict | None:
    progress_path = Path(PROGRESS_FILE)
    if not progress_path.exists():
        return None
    with progress_path.open("r", encoding="utf-8") as progress_file:
        return json.load(progress_file)


def save_progress(progress: dict) -> None:
    with Path(PROGRESS_FILE).open("w", encoding="utf-8") as progress_file:
        json.dump(progress, progress_file, ensure_ascii=False, indent=2)


class WeeklyScanWorkbook:
    """Gere o Excel datado preservando formatação de cores."""

    def __init__(self, path: str) -> None:
        self.path = Path(path)
        self.workbook = load_workbook(self.path)
        self.worksheet = self.workbook["Produtos"]
        self.url_column = self._find_url_column()
        self.url_to_row: dict[str, int] = {}
        self._rebuild_url_index()

    def _find_url_column(self) -> int:
        for column_index in range(1, self.worksheet.max_column + 1):
            header = self.worksheet.cell(1, column_index).value
            if header == "URL":
                return column_index
        raise SystemExit("Coluna 'URL' não encontrada no Excel.")

    def _rebuild_url_index(self) -> None:
        self.url_to_row.clear()
        for row_index in range(2, self.worksheet.max_row + 1):
            raw_url = self.worksheet.cell(row_index, self.url_column).value
            if not raw_url:
                continue
            self.url_to_row[normalize_url(str(raw_url))] = row_index

    def mark_removed(self, url: str) -> bool:
        row_index = self.url_to_row.get(normalize_url(url))
        if not row_index:
            return False
        for column_index in range(1, self.worksheet.max_column + 1):
            self.worksheet.cell(row_index, column_index).font = RED_FONT
        return True

    def append_record(self, record: dict[str, str]) -> None:
        row_index = self.worksheet.max_row + 1
        for column_index, column_name in enumerate(auchan.EXCEL_COLUMNS, start=1):
            cell = self.worksheet.cell(row_index, column_index, record.get(column_name, ""))
            cell.font = BLUE_FONT
        self.url_to_row[normalize_url(record["URL"])] = row_index

    def save(self) -> None:
        self.workbook.save(self.path)


def seed_progress_from_excel(progress: dict, output_file: str) -> None:
    excel_path = Path(output_file)
    if not excel_path.exists():
        raise SystemExit(f"Ficheiro de saída não encontrado: {output_file}")

    dataframe = pd.read_excel(excel_path)
    missing_columns = [column for column in auchan.EXCEL_COLUMNS if column not in dataframe.columns]
    if missing_columns:
        raise SystemExit(f"Colunas em falta no Excel: {missing_columns}")

    records = dataframe[auchan.EXCEL_COLUMNS].astype(str).to_dict(orient="records")
    baseline_urls = sorted(
        {
            normalize_url(url)
            for url in dataframe["URL"].dropna().astype(str).tolist()
            if url and url not in ("N/A", "nan")
        }
    )

    progress["records"] = records
    progress["baseline_urls"] = baseline_urls
    save_progress(progress)
    log(
        f"Baseline carregado — {len(records)} produtos, "
        f"{len(baseline_urls)} URLs únicas em '{output_file}'."
    )


def commit_checkpoint_to_repo(progress: dict, *, reason: str) -> bool:
    if not _git_checkpoint_enabled:
        return False

    output_file = progress.get("output_file", "")
    files_to_add = [name for name in (PROGRESS_FILE, output_file) if (REPO_ROOT / name).exists()]
    if not files_to_add:
        return False

    scan_date = progress.get("scan_date", "desconhecida")
    new_count = progress.get("new_products_added", 0)
    total_count = len(progress.get("records", []))
    message = (
        f"Checkpoint weekly scan Auchan {scan_date}: {new_count} novos, "
        f"{total_count} total — {reason}"
    )

    try:
        subprocess.run(
            ["git", "add", "--", *files_to_add],
            cwd=REPO_ROOT,
            capture_output=True,
            text=True,
            check=True,
        )
        commit_result = subprocess.run(
            ["git", "commit", "-m", message],
            cwd=REPO_ROOT,
            capture_output=True,
            text=True,
            check=False,
        )
        if commit_result.returncode != 0:
            combined = f"{commit_result.stdout}\n{commit_result.stderr}".lower()
            if "nothing to commit" in combined:
                return False
            log(f"[Git] git commit falhou: {commit_result.stderr.strip()}")
            return False

        branch_name = subprocess.check_output(
            ["git", "rev-parse", "--abbrev-ref", "HEAD"],
            cwd=REPO_ROOT,
            text=True,
        ).strip()

        for attempt, wait_seconds in enumerate(GIT_PUSH_BACKOFF_SECONDS, start=1):
            push_result = subprocess.run(
                ["git", "push", "-u", "origin", branch_name],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
                check=False,
            )
            if push_result.returncode == 0:
                log(f"[Git] Checkpoint enviado para origin/{branch_name} ({reason}).")
                return True
            if attempt == len(GIT_PUSH_BACKOFF_SECONDS):
                log(f"[Git] push falhou: {push_result.stderr.strip()}")
                return False
            log(f"[Git] push falhou, nova tentativa em {wait_seconds}s...")
            time.sleep(wait_seconds)
    except Exception as error:
        log(f"[Git] Erro ao commitar checkpoint: {error}")
        return False

    return False


def maybe_checkpoint(progress: dict, workbook: WeeklyScanWorkbook, *, use_git: bool, reason: str) -> None:
    workbook.save()
    save_progress(progress)
    new_count = progress.get("new_products_added", 0)
    log(
        f"[Checkpoint] {new_count} produtos novos processados | "
        f"{len(progress.get('records', []))} total em '{progress['output_file']}'."
    )
    if use_git:
        commit_checkpoint_to_repo(progress, reason=reason)


def prepare_output_file(source_file: str, output_file: str) -> None:
    source_path = Path(source_file)
    output_path = Path(output_file)
    if not source_path.exists():
        raise SystemExit(f"Ficheiro fonte não encontrado: {source_file}")
    if not output_path.exists():
        shutil.copy2(source_path, output_path)
        log(f"Cópia criada: '{source_file}' → '{output_file}'.")
    else:
        log(f"Retomando ficheiro existente: '{output_file}'.")


def initialize_progress(source_file: str, output_file: str, scan_date: str) -> dict:
    prepare_output_file(source_file, output_file)
    progress = empty_progress(source_file, output_file, scan_date)
    seed_progress_from_excel(progress, output_file)
    return progress


def scan_category_listings(
    page: Page,
    target: auchan.CategoryTarget,
    progress: dict,
    *,
    max_pages: int | None = None,
    max_products: int | None = None,
) -> None:
    listing_key = category_key(target.listing_url)
    category_name = target.display_name

    if listing_key in progress.get("completed_categories", []):
        log(f"[Varredura: {category_name}] Já concluída — a saltar.")
        return

    log(f"\n{'=' * 70}\n[Varredura: {category_name}]\nURL: {target.listing_url}\n{'=' * 70}")

    products = auchan.collect_all_category_products(
        page,
        target,
        navigate_via_menu_flag=False,
        max_pages=max_pages,
        max_products=max_products,
        progress=progress,
    )

    site_products: dict[str, dict[str, str]] = dict(progress.get("site_products", {}))
    site_urls = set(progress.get("site_urls", []))

    for product in products:
        url = normalize_url(product.get("url", ""))
        if not url:
            continue
        site_urls.add(url)
        site_products[url] = {
            "pid": str(product.get("pid", "")),
            "name": str(product.get("name", "")),
            "url": url,
            "unit_price": str(product.get("unit_price", "")),
            "list_price": str(product.get("list_price", "")),
            "sales_price": str(product.get("sales_price", "")),
        }

    completed = list(progress.get("completed_categories", []))
    if listing_key not in completed:
        completed.append(listing_key)
    progress["completed_categories"] = completed
    progress["site_urls"] = sorted(site_urls)
    progress["site_products"] = site_products
    progress["in_progress"] = None
    save_progress(progress)

    log(
        f"[Varredura: {category_name}] Concluída — "
        f"{len(products)} produtos na listagem | {len(site_urls)} URLs únicas no site."
    )


def compute_differences(progress: dict, *, total_categories: int) -> None:
    baseline_urls = {normalize_url(url) for url in progress.get("baseline_urls", [])}
    site_urls = {normalize_url(url) for url in progress.get("site_urls", [])}

    new_urls = sorted(site_urls - baseline_urls)
    all_categories_scanned = len(progress.get("completed_categories", [])) >= total_categories
    removed_urls = sorted(baseline_urls - site_urls) if all_categories_scanned else []

    progress["new_urls_pending"] = new_urls
    progress["removed_urls"] = removed_urls
    save_progress(progress)

    log(
        f"\n[Comparação] Baseline: {len(baseline_urls)} | Site: {len(site_urls)} | "
        f"Novos: {len(new_urls)} | Removidos: {len(removed_urls)}"
    )
    if not all_categories_scanned:
        log("[Comparação] Varredura parcial — removidos só são calculados com todas as secções.")


def extract_new_products(
    page: Page,
    progress: dict,
    workbook: WeeklyScanWorkbook,
    *,
    use_git: bool,
    max_new_products: int | None = None,
) -> None:
    pending_urls = [
        url
        for url in progress.get("new_urls_pending", [])
        if url not in set(progress.get("new_urls_processed", []))
    ]
    if max_new_products is not None:
        pending_urls = pending_urls[:max_new_products]

    site_products: dict[str, dict[str, str]] = progress.get("site_products", {})
    records: list[dict[str, str]] = progress.setdefault("records", [])
    processed = set(progress.get("new_urls_processed", []))
    total = len(pending_urls)

    if not pending_urls:
        log("[Novos] Nenhum produto novo por processar.")
        return

    log(f"\n[Novos] A extrair {total} produto(s) novo(s).")

    for index, url in enumerate(pending_urls, start=1):
        product_stub = site_products.get(url, {"url": url, "name": "", "pid": ""})
        category_name = "Varredura Semanal"

        auchan.log_product_progress(
            category_name,
            index,
            total,
            status="processing",
            product_name=str(product_stub.get("name", "")).strip(),
        )

        try:
            product_data = auchan.extract_product_details(page, product_stub, category_name)
        except Exception as error:
            auchan.log_product_progress(
                category_name,
                index,
                total,
                status="error",
                details=str(error),
            )
            product_data = auchan.empty_product_record(product_stub, category_name)

        records.append(product_data)
        workbook.append_record(product_data)
        processed.add(url)
        progress["records"] = records
        progress["new_urls_processed"] = sorted(processed)
        progress["new_products_added"] = int(progress.get("new_products_added", 0)) + 1
        save_progress(progress)

        price_info = product_data["Preço Regular"]
        if product_data.get("Preço Promocional") not in ("N/A", ""):
            price_info = f"{product_data['Preço Regular']} (promo: {product_data['Preço Promocional']})"

        auchan.log_product_progress(
            category_name,
            index,
            total,
            status="done",
            product_name=product_data["Descrição do Produto"],
            details=(
                f"{product_data['Caminho Categorias']} "
                f"| EAN: {product_data['EAN / Referência']} "
                f"| Preço: {price_info}"
            ),
        )

        if progress["new_products_added"] % CHECKPOINT_EVERY == 0:
            maybe_checkpoint(
                progress,
                workbook,
                use_git=use_git,
                reason=f"{progress['new_products_added']} novos",
            )

        if index < total:
            auchan.human_delay(0.8, 1.8)


def mark_removed_products(progress: dict, workbook: WeeklyScanWorkbook) -> None:
    removed_urls = progress.get("removed_urls", [])
    already_marked = set(progress.get("removed_urls_marked", []))
    pending = [url for url in removed_urls if url not in already_marked]

    if not pending:
        log("[Removidos] Nenhum produto por marcar a vermelho.")
        return

    log(f"\n[Removidos] A marcar {len(pending)} produto(s) a vermelho.")
    marked = list(already_marked)

    for index, url in enumerate(pending, start=1):
        if workbook.mark_removed(url):
            log(f"[Removidos] {index}/{len(pending)} marcado | {url}")
        else:
            log(f"[Removidos] {index}/{len(pending)} URL não encontrada no Excel | {url}")
        marked.append(url)

    progress["removed_urls_marked"] = sorted(set(marked))
    save_progress(progress)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Varredura semanal Auchan — novos (azul) e removidos (vermelho)."
    )
    parser.add_argument("--test", action="store_true", help="2 categorias, até 2 produtos novos.")
    parser.add_argument("--max-categories", type=int, help="Limita categorias na varredura.")
    parser.add_argument("--max-pages", type=int, help="Limita páginas por categoria.")
    parser.add_argument("--max-new-products", type=int, help="Limita produtos novos a extrair.")
    parser.add_argument("--date", help="Data YYYY-MM-DD para o ficheiro (predefinido: hoje).")
    parser.add_argument("--source", default=SOURCE_FILE, help=f"Excel base (predefinido: {SOURCE_FILE}).")
    parser.add_argument("--reset-progress", action="store_true", help="Apaga progresso da varredura.")
    parser.add_argument("--no-git", action="store_true", help="Desativa commit/push automático.")
    return parser.parse_args()


def main() -> None:
    global _git_checkpoint_enabled

    args = parse_args()
    scan_date = args.date or date.today().isoformat()
    output_file = output_filename_for_date(scan_date)
    use_git = not args.no_git
    _git_checkpoint_enabled = use_git

    max_categories = args.max_categories if args.max_categories else (2 if args.test else None)
    max_new_products = args.max_new_products if args.max_new_products else (2 if args.test else None)

    if args.reset_progress and Path(PROGRESS_FILE).exists():
        Path(PROGRESS_FILE).unlink()
        log(f"Progresso removido: {PROGRESS_FILE}")

    progress = load_progress()
    if progress and progress.get("scan_date") != scan_date:
        log(
            f"Progresso existente é de {progress.get('scan_date')} — "
            f"a iniciar nova varredura para {scan_date}."
        )
        progress = None

    if progress is None:
        progress = initialize_progress(args.source, output_file, scan_date)
    elif not progress.get("baseline_urls"):
        seed_progress_from_excel(progress, output_file)

    progress["output_file"] = output_file
    workbook = WeeklyScanWorkbook(output_file)
    targets = build_weekly_scan_targets()

    log(f"Weekly scan Auchan — data {scan_date}")
    log(f"Fonte: {args.source} | Saída: {output_file}")
    log(f"Secções a varrer: {len(targets)}")
    log(f"Baseline: {len(progress.get('baseline_urls', []))} URLs")
    if use_git:
        log(f"Git: commit/push a cada {CHECKPOINT_EVERY} produtos novos.")

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        context = browser.new_context(
            locale="pt-PT",
            viewport={"width": 1920, "height": 1080},
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/120.0.0.0 Safari/537.36"
            ),
        )
        page = context.new_page()

        try:
            if progress.get("phase", "scanning") == "scanning":
                pending_targets = [
                    target
                    for target in targets
                    if category_key(target.listing_url) not in progress.get("completed_categories", [])
                ]
                if max_categories:
                    pending_targets = pending_targets[:max_categories]

                log(f"\n[Fase 1] Varredura de listagens — {len(pending_targets)} categorias pendentes.")
                for target in pending_targets:
                    scan_category_listings(
                        page,
                        target,
                        progress,
                        max_pages=args.max_pages,
                        max_products=2 if args.test else None,
                    )
                    if not args.test:
                        auchan.human_delay(1.0, 2.0)

                compute_differences(progress, total_categories=len(targets))
                progress["phase"] = "extracting_new"
                save_progress(progress)

            if progress.get("phase") == "extracting_new":
                log("\n[Fase 2] Extração de produtos novos.")
                extract_new_products(
                    page,
                    progress,
                    workbook,
                    use_git=use_git,
                    max_new_products=max_new_products,
                )
                progress["phase"] = "marking_removed"
                save_progress(progress)

            if progress.get("phase") == "marking_removed":
                log("\n[Fase 3] Marcação de produtos removidos.")
                mark_removed_products(progress, workbook)
                progress["phase"] = "completed"
                save_progress(progress)

            workbook.save()
            maybe_checkpoint(progress, workbook, use_git=use_git, reason="varredura concluída")

            log(
                f"\nConcluído. Ficheiro: '{output_file}' | "
                f"{len(progress.get('records', []))} produtos | "
                f"{progress.get('new_products_added', 0)} novos | "
                f"{len(progress.get('removed_urls_marked', []))} marcados a vermelho."
            )
        finally:
            context.close()
            browser.close()


if __name__ == "__main__":
    main()
