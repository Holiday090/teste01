"""
Preenche as colunas vazias do Promo_11_08.xlsx com dados do Continente Online.

Colunas preenchidas:
  - Encontrado?  -> SIM / NÃO
  - PROMO?       -> SIM / NÃO
  - Preço Normal -> preço regular (ou PVPR quando em promoção)
  - Preço Promo  -> preço promocional (vazio se não houver promoção)
"""

from __future__ import annotations

import argparse
import asyncio
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from playwright.async_api import Browser, Page, async_playwright

BASE_URL = "https://www.continente.pt/"
SEARCH_URL = f"{BASE_URL}pesquisa/?q={{ean}}"
INPUT_FILE = Path("Promo_11_08.xlsx")
PROGRESS_FILE = Path("continente_promo_progress.json")

MAX_RETRIES = 3
REQUEST_DELAY_SECONDS = 0.8
PAGE_TIMEOUT_MS = 30_000


def normalizar_ean(valor) -> str:
    if valor is None:
        return ""
    texto = str(valor).strip().split(".")[0]
    return texto


def parse_preco(texto: str) -> float | None:
    if not texto:
        return None
    limpo = (
        texto.replace("\xa0", " ")
        .replace("€", "")
        .replace("PVPR", "")
        .replace("/un", "")
        .strip()
    )
    limpo = re.sub(r"\s+", "", limpo)
    match = re.search(r"(\d+)[,.](\d{2})", limpo)
    if match:
        return float(f"{match.group(1)}.{match.group(2)}")
    match = re.search(r"(\d+)", limpo)
    if match:
        return float(match.group(1))
    return None


def resultado_nao_encontrado(ean: str) -> dict:
    return {
        "ean": ean,
        "encontrado": "NÃO",
        "promo": None,
        "preco_normal": None,
        "preco_promo": None,
    }


async def aceitar_cookies(page: Page) -> None:
    for seletor in ("#onetrust-accept-btn-handler", "button:has-text('Aceitar')"):
        botao = page.locator(seletor).first
        if await botao.count() > 0 and await botao.is_visible():
            await botao.click()
            await page.wait_for_timeout(400)
            return


async def extrair_dados_tile(tile) -> dict:
    nome = ""
    nome_loc = tile.locator("h2.pwc-tile--description").first
    if await nome_loc.count() > 0:
        nome = (await nome_loc.inner_text()).strip()

    link_loc = tile.locator("a[href*='/produto/']").first
    url_produto = ""
    if await link_loc.count() > 0:
        url_produto = (await link_loc.get_attribute("href")) or ""

    preco_atual_texto = ""
    preco_atual_loc = tile.locator(".pwc-tile--price-primary").first
    if await preco_atual_loc.count() > 0:
        preco_atual_texto = (await preco_atual_loc.inner_text()).strip()

    preco_anterior_texto = ""
    list_loc = tile.locator(".prices-wrapper .list").first
    if await list_loc.count() > 0:
        preco_anterior_texto = (await list_loc.inner_text()).strip()

    em_promocao = False
    if preco_anterior_texto and parse_preco(preco_anterior_texto) is not None:
        em_promocao = True

    badge_promo = tile.locator(
        ".ct-product-tile-badge--promotional, .ct-product-tile-badge--pvpr"
    ).first
    if await badge_promo.count() > 0:
        em_promocao = True

    preco_atual = parse_preco(preco_atual_texto)
    preco_anterior = parse_preco(preco_anterior_texto)

    if em_promocao and preco_anterior and preco_atual and preco_anterior <= preco_atual:
        em_promocao = False
        preco_anterior = None

    if em_promocao:
        return {
            "nome_produto": nome,
            "em_promocao": "SIM",
            "preco_normal": preco_anterior,
            "preco_promo": preco_atual,
            "url_produto": url_produto,
        }

    return {
        "nome_produto": nome,
        "em_promocao": "NÃO",
        "preco_normal": preco_atual,
        "preco_promo": None,
        "url_produto": url_produto,
    }


async def pesquisar_ean(page: Page, ean: str) -> dict:
    ultimo_erro: Exception | None = None

    for tentativa in range(1, MAX_RETRIES + 1):
        try:
            await page.goto(
                SEARCH_URL.format(ean=ean),
                wait_until="domcontentloaded",
                timeout=PAGE_TIMEOUT_MS,
            )
            await page.wait_for_load_state("networkidle", timeout=PAGE_TIMEOUT_MS)

            body = await page.locator("body").inner_text()
            for linha in body.split("\n"):
                ll = linha.lower()
                if "encontr" in ll and "produto" in ll and "0 produto" in ll:
                    return resultado_nao_encontrado(ean)

            await page.wait_for_selector(".js-product-grid, .product-tile", timeout=8_000)

            tiles = page.locator(".product-tile")
            total = await tiles.count()
            if total == 0:
                return resultado_nao_encontrado(ean)

            tile = tiles.first
            if total > 1:
                for i in range(total):
                    candidato = tiles.nth(i)
                    texto = (await candidato.inner_text()).replace(" ", "")
                    if ean in texto or ean.lstrip("0") in texto:
                        tile = candidato
                        break

            dados = await extrair_dados_tile(tile)
            return {
                "ean": ean,
                "encontrado": "SIM",
                "promo": dados["em_promocao"],
                "preco_normal": dados["preco_normal"],
                "preco_promo": dados["preco_promo"],
            }

        except Exception as exc:
            ultimo_erro = exc
            if tentativa < MAX_RETRIES:
                await asyncio.sleep(1.5 * tentativa)

    print(f"  ERRO em {ean} após {MAX_RETRIES} tentativas: {ultimo_erro}", file=sys.stderr)
    return resultado_nao_encontrado(ean)


def carregar_progresso() -> dict[str, dict]:
    if PROGRESS_FILE.exists():
        return json.loads(PROGRESS_FILE.read_text(encoding="utf-8"))
    return {}


def guardar_progresso(progresso: dict[str, dict]) -> None:
    PROGRESS_FILE.write_text(json.dumps(progresso, ensure_ascii=False, indent=2), encoding="utf-8")


def ler_eans_do_excel(caminho: Path) -> list[tuple[int, str]]:
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ws = wb["Folha1"]
    linhas: list[tuple[int, str]] = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        ean = normalizar_ean(row[0])
        if ean:
            linhas.append((idx, ean))
    wb.close()
    return linhas


def escrever_resultados_no_excel(caminho: Path, resultados_por_ean: dict[str, dict]) -> None:
    wb = openpyxl.load_workbook(caminho)
    ws = wb["Folha1"]

    for linha in range(2, ws.max_row + 1):
        ean = normalizar_ean(ws.cell(linha, 1).value)
        if not ean or ean not in resultados_por_ean:
            continue
        dados = resultados_por_ean[ean]
        ws.cell(linha, 2, dados["encontrado"])
        ws.cell(linha, 3, dados["promo"] if dados["encontrado"] == "SIM" else None)
        ws.cell(linha, 4, dados["preco_normal"])
        ws.cell(linha, 5, dados["preco_promo"])

    wb.save(caminho)
    wb.close()


def formatar_progresso(dados: dict) -> str:
    if dados["encontrado"] == "NÃO":
        return "NÃO ENCONTRADO"
    if dados["promo"] == "SIM":
        anterior = dados["preco_normal"]
        atual = dados["preco_promo"]
        return f"ENCONTRADO | PROMOÇÃO | {anterior:.2f} € → {atual:.2f} €"
    preco = dados["preco_normal"]
    if preco is not None:
        return f"ENCONTRADO | SEM PROMOÇÃO | {preco:.2f} €"
    return "ENCONTRADO | SEM PROMOÇÃO"


async def processar_eans(
    eans: list[str],
    headless: bool = True,
    limite: int | None = None,
) -> dict[str, dict]:
    progresso = carregar_progresso()
    pendentes = [ean for ean in eans if ean not in progresso]
    if limite is not None:
        pendentes = pendentes[:limite]

    if not pendentes:
        return progresso

    async with async_playwright() as playwright:
        browser: Browser = await playwright.chromium.launch(headless=headless)
        context = await browser.new_context(locale="pt-PT")
        page = await context.new_page()

        await page.goto(BASE_URL, wait_until="domcontentloaded", timeout=PAGE_TIMEOUT_MS)
        await aceitar_cookies(page)

        ja_processados = len(progresso)
        total = len(eans)
        for indice, ean in enumerate(pendentes, start=1):
            posicao = ja_processados + indice
            dados = await pesquisar_ean(page, ean)
            progresso[ean] = dados
            guardar_progresso(progresso)
            print(f"[{posicao}/{total}] {ean} → {formatar_progresso(dados)}", flush=True)
            await asyncio.sleep(REQUEST_DELAY_SECONDS)

        await browser.close()

    return progresso


async def main() -> None:
    parser = argparse.ArgumentParser(description="Preenche Promo_11_08.xlsx com dados do Continente")
    parser.add_argument("--input", default=str(INPUT_FILE))
    parser.add_argument("--limit", type=int, default=None, help="Processar apenas N EANs pendentes")
    parser.add_argument("--headed", action="store_true", help="Abrir browser visível")
    parser.add_argument("--reset", action="store_true", help="Apagar progresso e recomeçar")
    args = parser.parse_args()

    caminho = Path(args.input)
    if args.reset and PROGRESS_FILE.exists():
        PROGRESS_FILE.unlink()

    linhas = ler_eans_do_excel(caminho)
    eans = [ean for _, ean in linhas]
    print(f"A processar {len(eans)} EANs de {caminho.name}...", flush=True)

    progresso = await processar_eans(eans, headless=not args.headed, limite=args.limit)
    escrever_resultados_no_excel(caminho, progresso)

    encontrados = sum(1 for d in progresso.values() if d["encontrado"] == "SIM")
    print(f"\nConcluído: {encontrados}/{len(eans)} encontrados. Ficheiro atualizado: {caminho}")
    print(
        "Nota: o Excel foi gravado no disco local. Para ver a alteração no GitHub, "
        "é preciso fazer commit e push (git add, git commit, git push)."
    )


if __name__ == "__main__":
    asyncio.run(main())
